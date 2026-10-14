import csv
import io
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.response import Response

try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    openpyxl = None
    Font = None
    get_column_letter = None

SAO_PAULO_TZ = ZoneInfo("America/Sao_Paulo")

# Expressão para detectar injeção de fórmula em strings
FORMULA_PREFIXES = ("=", "+", "-", "@")


def sanitize_cell_value(value):
    """
    Neutraliza texto iniciado por =, +, - ou @ para impedir injeção de fórmula
    ao abrir no Excel ou outros leitores de planilhas.
    Valores numéricos e datas não são alterados.
    """
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith(FORMULA_PREFIXES):
            return f"'{value}"
    return value


def to_sao_paulo(dt_value):
    """
    Converte datetime ou date para o fuso horário America/Sao_Paulo.
    """
    if dt_value is None:
        return None
    if isinstance(dt_value, datetime):
        if timezone.is_aware(dt_value):
            return timezone.localtime(dt_value, SAO_PAULO_TZ)
        return dt_value.replace(tzinfo=SAO_PAULO_TZ)
    return dt_value


class ExportColumn:
    """
    Representação de uma coluna para exportação:
    - header: Título da coluna
    - key: chave ou função extratora
    - data_type: 'text', 'currency', 'datetime', 'date', 'number'
    """

    def __init__(self, header, key=None, data_type="text"):
        self.header = header
        self.key = key or header
        self.data_type = data_type


def export_data_response(
    filename_base, headers, rows, file_format="csv", column_types=None
):
    """
    Gera HttpResponse contendo arquivo CSV ou XLSX com base no file_format.

    :param filename_base: Nome base do arquivo (sem extensão)
    :param headers: Lista de strings com os cabeçalhos das colunas
    :param rows: Lista de listas/tuplas com os valores das células
    :param file_format: 'csv' (padrão) ou 'xlsx'
    :param column_types: Lista opcional de tipos ('text', 'currency', 'datetime', 'date', 'number')
                         correspondente a cada coluna
    :return: HttpResponse ou Response(400)
    """
    norm_format = (file_format or "csv").strip().lower()

    if norm_format not in ("csv", "xlsx"):
        return Response(
            {"detail": "Formato inválido. Use 'csv' ou 'xlsx'."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if column_types is None:
        column_types = ["text"] * len(headers)

    if norm_format == "csv":
        return _generate_csv_response(filename_base, headers, rows, column_types)
    return _generate_xlsx_response(filename_base, headers, rows, column_types)


def _generate_csv_response(filename_base, headers, rows, column_types):
    """
    Gera CSV codificado em UTF-8 com BOM e delimitador ';'.
    """
    output = io.StringIO()
    # Adiciona o BOM do UTF-8 para o Excel abrir com acentuação correta
    output.write("\ufeff")

    writer = csv.writer(
        output, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n"
    )
    writer.writerow(headers)

    for row in rows:
        formatted_row = []
        for idx, val in enumerate(row):
            col_type = column_types[idx] if idx < len(column_types) else "text"
            formatted_val = _format_csv_cell(val, col_type)
            formatted_row.append(formatted_val)
        writer.writerow(formatted_row)

    csv_data = output.getvalue().encode("utf-8")
    response = HttpResponse(csv_data, content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename_base}.csv"'
    return response


def _format_csv_cell(val, col_type):
    if val is None:
        return ""

    if col_type == "currency":
        if isinstance(val, (int, float, Decimal)):
            # Formato monetário brasileiro: R$ 1.234,56
            return (
                f"R$ {val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            )
        return sanitize_cell_value(str(val))

    if col_type in ("datetime", "date"):
        sp_val = to_sao_paulo(val)
        if isinstance(sp_val, datetime):
            return sp_val.strftime("%d/%m/%Y %H:%M")
        if isinstance(sp_val, date):
            return sp_val.strftime("%d/%m/%Y")
        return sanitize_cell_value(str(val))

    if col_type == "number":
        if isinstance(val, (int, float, Decimal)):
            return str(val)
        return sanitize_cell_value(str(val))

    return sanitize_cell_value(str(val))


def _generate_xlsx_response(filename_base, headers, rows, column_types):
    """
    Gera planilha Excel (.xlsx) usando openpyxl:
    - Cabeçalho em negrito
    - Primeira linha congelada (freeze panes A2)
    - Largura de coluna auto-ajustada
    - Valores monetários nativos com máscara R$ #,##0.00
    - Datas nativas com máscara dd/mm/yyyy hh:mm
    """
    if openpyxl is None:  # pragma: no cover
        raise RuntimeError("openpyxl não está instalado.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"

    # Congela a primeira linha
    ws.freeze_panes = "A2"

    # Cabeçalho
    ws.append(headers)
    header_font = Font(bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font

    # Linhas de dados
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, val in enumerate(row, start=1):
            col_type = (
                column_types[col_idx - 1] if col_idx - 1 < len(column_types) else "text"
            )
            cell = ws.cell(row=row_idx, column=col_idx)
            _set_xlsx_cell(cell, val, col_type)

    # Ajuste automático de largura de colunas
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1] or ""))
        for row_idx in range(2, len(rows) + 2):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response = HttpResponse(output.getvalue(), content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename_base}.xlsx"'
    return response


def _set_xlsx_cell(cell, val, col_type):
    if val is None:
        cell.value = ""
        return

    if col_type == "currency":
        try:
            numeric_val = float(val)
            cell.value = numeric_val
            cell.number_format = '"R$"\\ #,##0.00'
        except (ValueError, TypeError):
            cell.value = sanitize_cell_value(str(val))
        return

    if col_type in ("datetime", "date"):
        sp_val = to_sao_paulo(val)
        if isinstance(sp_val, datetime):
            # Excel prefere datetime naive
            naive_dt = sp_val.replace(tzinfo=None)
            cell.value = naive_dt
            cell.number_format = "dd/mm/yyyy hh:mm"
        elif isinstance(sp_val, date):
            cell.value = sp_val
            cell.number_format = "dd/mm/yyyy"
        else:
            cell.value = sanitize_cell_value(str(val))
        return

    if col_type == "number":
        try:
            if isinstance(val, (int, float, Decimal)):
                cell.value = (
                    float(val) if isinstance(val, (float, Decimal)) else int(val)
                )
            else:
                cell.value = float(val) if "." in str(val) else int(val)
        except (ValueError, TypeError):
            cell.value = sanitize_cell_value(str(val))
        return

    # Texto
    cell.value = sanitize_cell_value(str(val))
