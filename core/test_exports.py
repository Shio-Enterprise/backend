import csv
import io
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

import openpyxl
from django.test import TestCase

from core.exports import (
    export_data_response,
    sanitize_cell_value,
    to_sao_paulo,
)


class CoreExportsTests(TestCase):
    def test_sanitize_cell_value_neutralizes_formula_prefixes(self):
        self.assertEqual(sanitize_cell_value("=SUM(A1:A10)"), "'=SUM(A1:A10)")
        self.assertEqual(
            sanitize_cell_value("+cmd|' /C calc'!A0"), "'+cmd|' /C calc'!A0"
        )
        self.assertEqual(sanitize_cell_value("-cmd"), "'-cmd")
        self.assertEqual(sanitize_cell_value("@HYPERLINK(...)"), "'@HYPERLINK(...)")
        self.assertEqual(sanitize_cell_value("  =1+1"), "'  =1+1")
        self.assertEqual(sanitize_cell_value("Texto Seguro"), "Texto Seguro")
        # Números não devem ser alterados
        self.assertEqual(sanitize_cell_value(-15), -15)
        self.assertEqual(sanitize_cell_value(Decimal("-20.50")), Decimal("-20.50"))

    def test_to_sao_paulo_converts_timezone(self):
        utc_dt = datetime(2026, 10, 14, 15, 30, tzinfo=ZoneInfo("UTC"))
        sp_dt = to_sao_paulo(utc_dt)
        self.assertEqual(sp_dt.hour, 12)
        self.assertEqual(sp_dt.minute, 30)

    def test_export_invalid_format_returns_400(self):
        response = export_data_response(
            filename_base="teste",
            headers=["Coluna 1"],
            rows=[["Valor"]],
            file_format="pdf",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Formato inválido", response.data["detail"])

    def test_export_csv_generation(self):
        headers = ["Nome", "Valor", "Status"]
        rows = [
            ["=Admin", Decimal("150.50"), "Ativo"],
            ["Cliente Comum", Decimal("0.00"), "Inativo"],
        ]
        column_types = ["text", "currency", "text"]

        response = export_data_response(
            filename_base="relatorio",
            headers=headers,
            rows=rows,
            file_format="csv",
            column_types=column_types,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn('filename="relatorio.csv"', response["Content-Disposition"])

        content = response.content
        # Verifica presença de UTF-8 BOM
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))

        # Decodifica com utf-8-sig (que consome o BOM) e valida linhas e delimitador ';'
        decoded = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(decoded), delimiter=";")
        lines = list(reader)
        self.assertEqual(lines[0], ["Nome", "Valor", "Status"])
        # Fórmula neutralizada no CSV
        self.assertEqual(lines[1][0], "'=Admin")
        self.assertEqual(lines[1][1], "R$ 150,50")

    def test_export_xlsx_generation(self):
        dt = datetime(2026, 10, 14, 18, 0, tzinfo=ZoneInfo("UTC"))
        headers = ["Produto", "Data", "Preço", "Fórmula"]
        rows = [
            ["Camiseta", dt, Decimal("99.90"), "+calc"],
        ]
        column_types = ["text", "datetime", "currency", "text"]

        response = export_data_response(
            filename_base="produtos",
            headers=headers,
            rows=rows,
            file_format="xlsx",
            column_types=column_types,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="produtos.xlsx"', response["Content-Disposition"])

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Cabeçalho em negrito
        self.assertTrue(ws.cell(row=1, column=1).font.bold)
        self.assertEqual(ws.cell(row=1, column=1).value, "Produto")

        # Freeze panes A2
        self.assertEqual(ws.freeze_panes, "A2")

        # Célula de valor e formatação de moeda
        price_cell = ws.cell(row=2, column=3)
        self.assertEqual(price_cell.value, 99.90)
        self.assertIn("R$", price_cell.number_format)

        # Célula de data
        date_cell = ws.cell(row=2, column=2)
        self.assertEqual(date_cell.number_format, "dd/mm/yyyy hh:mm")
        self.assertEqual(
            date_cell.value.hour, 15
        )  # 18:00 UTC -> 15:00 America/Sao_Paulo

        # Fórmula neutralizada
        formula_cell = ws.cell(row=2, column=4)
        self.assertEqual(formula_cell.value, "'+calc")
