import csv
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from authentication.admin_permissions import ADMIN_PERMISSION_CODES
from authentication.models import UserProfile, UserRole
from orders.models import (
    Coupon,
    CouponDiscountType,
    CustomerOrder,
    OrderItem,
    OrderStatus,
    Payment,
    PaymentMethod,
    PaymentStatus,
)
from products.models import Product, ProductVariation

User = get_user_model()


class AdminOrderExportTests(APITestCase):
    def setUp(self):
        self.url = "/api/orders/admin/export/"

        # Admin com permissão manage_orders
        self.admin = User.objects.create_user(
            email="admin.orders@example.com", name="Admin Pedidos", is_staff=True
        )
        UserProfile.objects.create(user=self.admin, role=UserRole.ADMIN)

        # Admin sem permissão manage_orders (só manage_customers)
        self.unauthorized_admin = User.objects.create_user(
            email="other.admin@example.com", name="Admin Clientes", is_staff=True
        )
        UserProfile.objects.create(user=self.unauthorized_admin, role=UserRole.ADMIN)

        # Cliente normal
        self.customer = User.objects.create_user(
            email="cliente@example.com", name="=Cliente Formula"
        )
        UserProfile.objects.create(user=self.customer, role=UserRole.CUSTOMER)

        self.permissions = {
            perm.codename: perm
            for perm in Permission.objects.filter(
                content_type__app_label="authentication",
                codename__in=ADMIN_PERMISSION_CODES,
            )
        }
        self.admin.user_permissions.set([self.permissions["manage_orders"]])
        self.unauthorized_admin.user_permissions.set(
            [self.permissions["manage_customers"]]
        )

        # Cria produto e variação
        self.product = Product.objects.create(
            name="Camiseta Minimalista",
            description="100% Algodão",
            base_price=Decimal("120.00"),
        )
        self.variation = ProductVariation.objects.create(
            product=self.product,
            size="M",
            color="Preto",
            sku="CAM-MIN-M-BLK",
            stock_quantity=10,
        )

        # Cupom
        self.coupon = Coupon.objects.create(
            code="DESC10",
            description="10% off",
            discount_type=CouponDiscountType.PERCENTAGE,
            discount_value=Decimal("10.00"),
        )

        # Pedido 1 (PAID)
        self.order_paid = CustomerOrder.objects.create(
            user=self.customer,
            status=OrderStatus.PAID,
            subtotal=Decimal("120.00"),
            discount_amount=Decimal("12.00"),
            coupon=self.coupon,
            shipping_cost=Decimal("15.00"),
            total_amount=Decimal("123.00"),
            shipping_zip_code="01001-000",
            shipping_street="Praça da Sé",
            shipping_number="1",
            shipping_neighborhood="Sé",
            shipping_city="São Paulo",
            shipping_state="SP",
        )
        OrderItem.objects.create(
            order=self.order_paid,
            variation=self.variation,
            # Formato salvo pelo checkout: nome do produto + variação
            product_name=f"{self.product.name} - M / Preto",
            sku_snapshot=self.variation.sku,
            quantity=1,
            unit_price=Decimal("120.00"),
        )
        Payment.objects.create(
            order=self.order_paid,
            method=PaymentMethod.PIX,
            status=PaymentStatus.PAID,
            total_amount=Decimal("123.00"),
            paid_at=timezone.now(),
        )

        # Pedido 2 (AWAITING_PAYMENT)
        self.customer_two = User.objects.create_user(
            email="outro@example.com", name="Outro Cliente"
        )
        UserProfile.objects.create(user=self.customer_two, role=UserRole.CUSTOMER)

        self.order_pending = CustomerOrder.objects.create(
            user=self.customer_two,
            status=OrderStatus.AWAITING_PAYMENT,
            subtotal=Decimal("240.00"),
            discount_amount=Decimal("0.00"),
            shipping_cost=Decimal("20.00"),
            total_amount=Decimal("260.00"),
            shipping_zip_code="20000-000",
            shipping_street="Rua da Carioca",
            shipping_number="10",
            shipping_neighborhood="Centro",
            shipping_city="Rio de Janeiro",
            shipping_state="RJ",
        )
        OrderItem.objects.create(
            order=self.order_pending,
            variation=self.variation,
            product_name=self.product.name,
            sku_snapshot=self.variation.sku,
            quantity=2,
            unit_price=Decimal("120.00"),
        )

    def test_anonymous_user_returns_401(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_unauthorized_admin_returns_403(self):
        self.client.force_authenticate(self.unauthorized_admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_authorized_admin_can_export_csv_by_default(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn('filename="pedidos.csv"', response["Content-Disposition"])

        content = response.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content), delimiter=";")
        lines = list(reader)

        headers = lines[0]
        self.assertEqual(
            headers,
            [
                "Número",
                "Data",
                "Cliente",
                "E-mail",
                "Itens",
                "Subtotal",
                "Desconto",
                "Cupom",
                "Frete",
                "Total",
                "Status do Pedido",
                "Status do Pagamento",
            ],
        )

        # 2 pedidos + 1 cabeçalho = 3 linhas
        self.assertEqual(len(lines), 3)

        # Cliente '=Cliente Formula' neutralizado
        self.assertTrue(
            any(
                row[2] == "''=Cliente Formula" or row[2] == "'=Cliente Formula"
                for row in lines[1:]
            )
        )

    def test_export_invalid_format_returns_400(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?file_format=xml")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_export_xlsx_format(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?file_format=xlsx")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.cell(row=1, column=1).font.bold)
        self.assertEqual(ws.max_row, 3)

    def test_filter_by_status(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?status=PAID&file_format=csv")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        content = response.content.decode("utf-8-sig")
        lines = list(csv.reader(io.StringIO(content), delimiter=";"))
        # Apenas 1 pedido + cabeçalho
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1][0], str(self.order_paid.id))

    def _csv_lines(self, response):
        content = response.content.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(content), delimiter=";"))

    def test_filter_by_search(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?search=outro cliente&file_format=csv")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        lines = self._csv_lines(response)
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1][0], str(self.order_pending.id))

    def test_search_ignores_email_when_customer_has_name(self):
        # A tela de Pedidos não busca por e-mail quando o cliente tem nome
        self.client.force_authenticate(self.admin)
        response = self.client.get(
            f"{self.url}?search=outro@example.com&file_format=csv"
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(self._csv_lines(response)), 1)

    def test_search_matches_email_when_customer_has_no_name(self):
        # Sem nome, a tela mostra o e-mail no lugar do nome
        self.customer_two.name = ""
        self.customer_two.save(update_fields=["name"])
        self.client.force_authenticate(self.admin)
        response = self.client.get(
            f"{self.url}?search=outro@example.com&file_format=csv"
        )
        lines = self._csv_lines(response)
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1][0], str(self.order_pending.id))

    def test_items_show_variation_once(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?status=PAID&file_format=csv")
        lines = self._csv_lines(response)
        self.assertEqual(lines[1][4], "1x Camiseta Minimalista (M / Preto)")

    def test_discount_and_coupon_in_separate_columns_csv(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?file_format=csv")
        rows = {row[0]: row for row in self._csv_lines(response)[1:]}

        paid = rows[str(self.order_paid.id)]
        self.assertEqual(paid[6], "R$ 12,00")
        self.assertEqual(paid[7], "DESC10")

        pending = rows[str(self.order_pending.id)]
        self.assertEqual(pending[6], "R$ 0,00")
        self.assertEqual(pending[7], "")

    def test_discount_is_numeric_currency_in_xlsx(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?status=PAID&file_format=xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active

        discount = ws.cell(row=2, column=7)
        self.assertEqual(discount.value, 12.0)
        self.assertEqual(discount.number_format, '"R$"\\ #,##0.00')
        self.assertEqual(ws.cell(row=2, column=8).value, "DESC10")
