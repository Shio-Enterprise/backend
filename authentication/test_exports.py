import csv
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from authentication.admin_permissions import ADMIN_PERMISSION_CODES
from authentication.models import NewsletterSubscriber, UserProfile, UserRole
from orders.models import (
    CustomerOrder,
    OrderStatus,
    Payment,
    PaymentMethod,
    PaymentStatus,
)

User = get_user_model()


class CustomerCRMExportTests(APITestCase):
    def setUp(self):
        self.url = "/api/auth/crm/customers/export/"

        # Admin com permissão manage_customers
        self.admin = User.objects.create_user(
            email="admin.crm@example.com", name="Admin CRM", is_staff=True
        )
        UserProfile.objects.create(user=self.admin, role=UserRole.ADMIN)

        # Admin sem permissão manage_customers
        self.unauthorized_admin = User.objects.create_user(
            email="admin.orders@example.com", name="Admin Orders", is_staff=True
        )
        UserProfile.objects.create(user=self.unauthorized_admin, role=UserRole.ADMIN)

        self.permissions = {
            perm.codename: perm
            for perm in Permission.objects.filter(
                content_type__app_label="authentication",
                codename__in=ADMIN_PERMISSION_CODES,
            )
        }
        self.admin.user_permissions.set([self.permissions["manage_customers"]])
        self.unauthorized_admin.user_permissions.set(
            [self.permissions["manage_orders"]]
        )

        # Cliente 1 - inscrito na newsletter e com pedidos
        self.customer1 = User.objects.create_user(
            email="ana@example.com", name="=Ana Paula"
        )
        UserProfile.objects.create(user=self.customer1, role=UserRole.CUSTOMER)
        NewsletterSubscriber.objects.create(
            email=self.customer1.email, consent_lgpd=True
        )

        # Pedido entregue para Cliente 1
        order1 = CustomerOrder.objects.create(
            user=self.customer1,
            status=OrderStatus.DELIVERED,
            subtotal=Decimal("150.00"),
            total_amount=Decimal("150.00"),
            shipping_zip_code="01001-000",
            shipping_street="Rua A",
            shipping_number="10",
            shipping_neighborhood="Centro",
            shipping_city="SP",
            shipping_state="SP",
        )
        Payment.objects.create(
            order=order1,
            method=PaymentMethod.PIX,
            status=PaymentStatus.PAID,
            total_amount=Decimal("150.00"),
            paid_at=timezone.now(),
        )

        # Cliente 2 - não inscrito na newsletter e sem pedidos
        self.customer2 = User.objects.create_user(
            email="bruno@example.com", name="Bruno Lima"
        )
        UserProfile.objects.create(user=self.customer2, role=UserRole.CUSTOMER)

    def test_anonymous_returns_401(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_unauthorized_admin_returns_403(self):
        self.client.force_authenticate(self.unauthorized_admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_authorized_admin_exports_csv_by_default(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn('filename="clientes.csv"', response["Content-Disposition"])

        content = response.content.decode("utf-8-sig")
        lines = list(csv.reader(io.StringIO(content), delimiter=";"))

        headers = lines[0]
        self.assertEqual(
            headers,
            [
                "Nome",
                "E-mail",
                "Data de Cadastro",
                "Pedidos Concluídos",
                "Total Gasto",
                "Última Compra",
                "Inscrito na Newsletter",
            ],
        )

        # 2 clientes + cabeçalho = 3 linhas
        self.assertEqual(len(lines), 3)

        # Verifica dados de Ana (inscrita) e fórmula neutralizada
        ana_row = next(r for r in lines[1:] if "ana@example.com" in r)
        self.assertEqual(ana_row[0], "'=Ana Paula")
        self.assertEqual(ana_row[3], "1")
        self.assertEqual(ana_row[4], "R$ 150,00")
        self.assertEqual(ana_row[6], "Sim")

        # Verifica dados de Bruno (não inscrito)
        bruno_row = next(r for r in lines[1:] if "bruno@example.com" in r)
        self.assertEqual(bruno_row[0], "Bruno Lima")
        self.assertEqual(bruno_row[3], "0")
        self.assertEqual(bruno_row[4], "R$ 0,00")
        self.assertEqual(bruno_row[6], "Não")

    def test_export_xlsx_format(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?file_format=xlsx")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.cell(row=1, column=1).font.bold)
        self.assertEqual(ws.max_row, 3)

    def test_search_filter(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?search=bruno&file_format=csv")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        content = response.content.decode("utf-8-sig")
        lines = list(csv.reader(io.StringIO(content), delimiter=";"))
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1][1], "bruno@example.com")
