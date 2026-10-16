import csv
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from rest_framework import status
from rest_framework.test import APITestCase

from authentication.admin_permissions import ADMIN_PERMISSION_CODES
from authentication.models import UserProfile, UserRole
from products.models import Category, DropCampaign, Product, ProductVariation

User = get_user_model()


class InventoryExportTests(APITestCase):
    def setUp(self):
        self.url = "/api/catalog/inventory/export/"

        # Admin com permissão manage_catalog
        self.admin = User.objects.create_user(
            email="admin.catalog@example.com", name="Admin Catálogo", is_staff=True
        )
        UserProfile.objects.create(user=self.admin, role=UserRole.ADMIN)

        # Admin sem permissão manage_catalog
        self.unauthorized_admin = User.objects.create_user(
            email="admin.orders@example.com", name="Admin Pedidos", is_staff=True
        )
        UserProfile.objects.create(user=self.unauthorized_admin, role=UserRole.ADMIN)

        self.permissions = {
            perm.codename: perm
            for perm in Permission.objects.filter(
                content_type__app_label="authentication",
                codename__in=ADMIN_PERMISSION_CODES,
            )
        }
        self.admin.user_permissions.set([self.permissions["manage_catalog"]])
        self.unauthorized_admin.user_permissions.set(
            [self.permissions["manage_orders"]]
        )

        # Drop e Categoria
        self.drop = DropCampaign.objects.create(
            name="Drop Inverno", slug="drop-inverno"
        )
        self.category = Category.objects.create(name="Moletons", slug="moletons")

        # Produto 1 (com drop e categoria)
        self.product1 = Product.objects.create(
            name="=Moletom Street",
            description="Moletom pesado",
            drop=self.drop,
            category=self.category,
            base_price=Decimal("250.00"),
            cost_price=Decimal("100.00"),
        )
        self.var1 = ProductVariation.objects.create(
            product=self.product1,
            size="G",
            color="Preto",
            sku="MOL-STR-G-BLK",
            stock_quantity=15,
        )

        # Produto 2 (sem drop nem categoria)
        self.product2 = Product.objects.create(
            name="Camiseta Básica",
            description="Algodão",
            base_price=Decimal("80.00"),
            cost_price=None,
        )
        self.var2 = ProductVariation.objects.create(
            product=self.product2,
            size="M",
            color="",
            sku="CAM-BAS-M",
            stock_quantity=5,
        )

    def test_anonymous_returns_401(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_unauthorized_admin_returns_403(self):
        self.client.force_authenticate(self.unauthorized_admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_authorized_admin_exports_csv_by_default(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn('filename="estoque.csv"', response["Content-Disposition"])

        content = response.content.decode("utf-8-sig")
        lines = list(csv.reader(io.StringIO(content), delimiter=";"))

        headers = lines[0]
        self.assertEqual(
            headers,
            [
                "Produto",
                "Drop",
                "Categoria",
                "Tamanho",
                "Cor",
                "Quantidade em Estoque",
                "Preço de Custo",
                "Preço de Venda",
            ],
        )

        # 2 variações + 1 cabeçalho = 3 linhas
        self.assertEqual(len(lines), 3)

        moletom_row = next(
            r for r in lines[1:] if "MOL-STR-G-BLK" in r or "Moletom" in r[0]
        )
        # Fórmula neutralizada no nome do produto
        self.assertEqual(moletom_row[0], "'=Moletom Street")
        self.assertEqual(moletom_row[1], "Drop Inverno")
        self.assertEqual(moletom_row[2], "Moletons")
        self.assertEqual(moletom_row[3], "G")
        self.assertEqual(moletom_row[4], "Preto")
        self.assertEqual(moletom_row[5], "15")
        self.assertEqual(moletom_row[6], "R$ 100,00")
        self.assertEqual(moletom_row[7], "R$ 250,00")

    def test_export_xlsx_format(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?file_format=xlsx")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.cell(row=1, column=1).font.bold)
        self.assertEqual(ws.max_row, 3)

    def test_filter_by_drop(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?drop={self.drop.id}&file_format=csv")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        content = response.content.decode("utf-8-sig")
        lines = list(csv.reader(io.StringIO(content), delimiter=";"))
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1][1], "Drop Inverno")

    def _csv_lines(self, response):
        content = response.content.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(content), delimiter=";"))

    def test_filter_by_search(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"{self.url}?search=camiseta&file_format=csv")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        lines = self._csv_lines(response)
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1][0], "Camiseta Básica")

    def test_search_only_by_product_name(self):
        # A tela de Produtos busca só pelo nome; SKU, tamanho e cor não entram
        self.client.force_authenticate(self.admin)
        for term in ("CAM-BAS-M", "G", "Preto"):
            response = self.client.get(f"{self.url}?search={term}&file_format=csv")
            self.assertEqual(len(self._csv_lines(response)), 1, term)

    def test_missing_cost_exported_as_empty_cell(self):
        self.client.force_authenticate(self.admin)

        response = self.client.get(f"{self.url}?search=camiseta&file_format=csv")
        self.assertEqual(self._csv_lines(response)[1][6], "")

        response = self.client.get(f"{self.url}?search=camiseta&file_format=xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
        self.assertIn(ws.cell(row=2, column=7).value, (None, ""))
